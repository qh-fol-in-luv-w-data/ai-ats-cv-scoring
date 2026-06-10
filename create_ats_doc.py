#!/usr/bin/env python3
"""
Tạo tài liệu ATS từ template: copy y chang cấu trúc XML, thay text + ảnh thực tế
"""
import copy, io, shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import lxml.etree as ET

# ── Paths ─────────────────────────────────────────────────────────────────────
TEMPLATE = Path('/Users/_qh.fol_/frappe-bench/apps/ai_ats/DAIT-Tài liệu mô tả AI Agent-Đánh Giá Thử Việc-2026.docx')
OUTPUT   = Path('/Users/_qh.fol_/frappe-bench/apps/ai_ats/DAIT-Tài-liệu-mô-tả-AI-Agent-Đánh-Giá-Ứng-Viên-2026.docx')
SS       = Path('/Users/_qh.fol_/.gemini/antigravity-ide/brain/c209b6ea-a7d9-4b64-a951-2459f413a1ef/screenshots')

# ── Load template và lấy các phần tử gốc ─────────────────────────────────────
tpl = Document(TEMPLATE)
body_tpl = tpl.element.body

# Lấy raw bytes ảnh từ template (image1.png = logo CT Group)
logo_rId = None
for rel in tpl.part.rels.values():
    if 'image' in rel.reltype and 'image1' in rel.target_ref:
        logo_rId = rel.rId
        logo_bytes = tpl.part.related_parts[rel.rId].blob
        break

print(f"Logo rId: {logo_rId}, size: {len(logo_bytes)//1024}KB")

# ── Tạo document mới từ template (kế thừa styles, page layout, header/footer) ─
shutil.copy(TEMPLATE, OUTPUT)
doc = Document(OUTPUT)
body = doc.element.body

# Xoá toàn bộ content (giữ sectPr)
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

# ── Helpers ───────────────────────────────────────────────────────────────────
NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
NS_PIC = 'http://schemas.openxmlformats.org/drawingml/2006/picture'

def add_image_para(doc, img_path, width_inches=5.5, caption=None, caption_italic=True):
    """Thêm ảnh + caption căn giữa"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.italic = caption_italic
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def set_cell_shading(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def table_borders(table, color='AAAAAA', sz='4'):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for bn in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{bn}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tblBorders.append(b)
    tblPr.append(tblBorders)

def make_tbl_width(table, total_cm=16):
    """Set table total width"""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(total_cm * 567)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

def para(doc, text='', bold=False, size=12, align='left', color=None, indent=0):
    p = doc.add_paragraph()
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right': p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)
    return p

def bullet(doc, text, size=12):
    p = doc.add_paragraph()
    p.add_run(text).font.size = Pt(size)
    # Thêm indent
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5)
    return p

def bold_inline(p, bold_text, normal_text='', size=12):
    r1 = p.add_run(bold_text)
    r1.bold = True; r1.font.size = Pt(size)
    if normal_text:
        r2 = p.add_run(normal_text)
        r2.font.size = Pt(size)

def add_toc(doc):
    """Thêm TOC field giống y template gốc"""
    # Copy SDT element (TOC) từ template
    sdt_orig = body_tpl[31]  # index 31 = sdt (TOC)
    sdt_copy = copy.deepcopy(sdt_orig)
    body.append(sdt_copy)

# ═══════════════════════════════════════════════════════════════════════════════
# TRANG BÌA — copy y chang structure từ template
# ═══════════════════════════════════════════════════════════════════════════════
# Copy elements [0..30] từ template, thay text cần thiết
cover_elements = []
for i in range(31):  # 0..30
    child = body_tpl[i]
    elem_copy = copy.deepcopy(child)
    cover_elements.append((i, elem_copy))

for i, elem in cover_elements:
    tag = elem.tag.split('}')[-1]
    
    if tag == 'p':
        text = ''.join(t.text or '' for t in elem.findall(f'.//{{{NS_W}}}t'))
        
        # Replace text trong paragraph
        replacements = {
            'AI ĐÁNH GIÁ HOÀN THÀNH THỬ VIỆC': 'AI ĐÁNH GIÁ ỨNG VIÊN',
            'Dự án: DAIT – AI Đánh Giá Hoàn Thành Thử Việc': 'Dự án: DAIT – AI Candidate Evaluation Persona',
            '01.6.2026': '04.6.2026',
        }
        for old, new in replacements.items():
            if old in text:
                for t_el in elem.findall(f'.//{{{NS_W}}}t'):
                    if t_el.text and old in t_el.text:
                        t_el.text = t_el.text.replace(old, new)
        
        body.append(elem)
        
    elif tag == 'tbl':
        # Bảng phiên bản (index 20) — copy và update ngày
        for t_el in elem.findall(f'.//{{{NS_W}}}t'):
            if t_el.text == '01.6.2026':
                t_el.text = '04.6.2026'
        body.append(elem)
        
    else:
        body.append(elem)

print("✅ Trang bìa done (copy từ template)")

# Empty para sau bìa + TOC
body.append(copy.deepcopy(body_tpl[30]))  # empty line

# ── TOC: copy SDT từ template ──────────────────────────────────────────────────
sdt_copy = copy.deepcopy(body_tpl[31])
# Thay chữ "MỤC LỤC" title nếu có
body.append(sdt_copy)
body.append(copy.deepcopy(body_tpl[32]))  # empty line sau TOC
body.append(copy.deepcopy(body_tpl[33]))  # empty line
print("✅ Mục lục done (copy TOC field từ template)")

# Page break trước nội dung
pb_p = OxmlElement('w:p')
pb_r = OxmlElement('w:r')
pb_br = OxmlElement('w:br')
pb_br.set(qn('w:type'), 'page')
pb_r.append(pb_br)
pb_p.append(pb_r)
body.append(pb_p)

# ═══════════════════════════════════════════════════════════════════════════════
# CHƯƠNG 1: GIỚI THIỆU HỆ THỐNG
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_heading('GIỚI THIỆU HỆ THỐNG', level=1)

doc.add_heading('1.1 Mục tiêu hệ thống', level=2)
doc.add_paragraph(
    'Hệ thống AI Đánh Giá Ứng Viên (AI Candidate Evaluation Persona) là nền tảng ứng dụng Trí tuệ '
    'Nhân tạo để tự động hóa và nâng cao chất lượng quy trình đánh giá, sàng lọc ứng viên tại CT Group '
    'theo triết lý "NoAI-NoHire". Hệ thống được xây dựng bởi Ban DAIT.'
)
doc.add_paragraph('Hệ thống hướng tới 3 mục tiêu cốt lõi sau:')

p = doc.add_paragraph(); bold_inline(p, 'Mục tiêu 1: Tự động hóa đánh giá và chấm điểm ứng viên')
bullet(doc, '- Chấm điểm tự động: Tự động scrape và phân tích bài Test AI (10 tiêu chí, thang 100), bài 5G Tiềm năng Quản lý (5 nhóm G1–G5, thang 100), bài IQ/EQ từ jobtest.vn và phỏng vấn Survey. Chấm SWAT Elite theo 4 trụ cột trọng số và đưa ra Quyết định ĐẠT/KHÔNG ĐẠT tự động.')

p = doc.add_paragraph(); bold_inline(p, 'Mục tiêu 2: Đảm bảo tính nhất quán và chuẩn hóa đánh giá')
bullet(doc, '- Đối chiếu đa nguồn: Tự động đối chiếu CV ↔ JD ↔ AI Test ↔ 5G Test ↔ Survey để đưa ra phân tích Strengths / Gaps / Best At có evidence trích dẫn cụ thể từng nguồn.')
bullet(doc, '- Áp dụng quy tắc tuyệt đối: Nếu AI Test label = "Non-AI" HOẶC 5G < 60 HOẶC SWAT < 6.0 → Decision = "KHÔNG ĐẠT" (theo văn hóa NoAI-NoHire của CT Group).')
bullet(doc, '- Chống hallucination: Mọi điểm số phải dựa trên dữ liệu thực từ bài test, không tự bịa khi dữ liệu rỗng.')

p = doc.add_paragraph(); bold_inline(p, 'Mục tiêu 3: Tối ưu hóa thời gian và nguồn lực cho bộ phận Tuyển dụng')
bullet(doc, '- Tiết kiệm thời gian HR: Thay thế 70–80% thời gian đọc và soát xét hồ sơ thủ công, giúp HR tập trung vào quyết định phỏng vấn thay vì đọc từng trang tài liệu.')
bullet(doc, '- Xuất báo cáo PDF tự động: Tạo báo cáo đánh giá định dạng A4 chuẩn ngay trong hệ thống, sẵn sàng trình Hiring Manager và Ban Giám đốc trong 30–60 giây.')

doc.add_heading('1.2 Đối tượng sử dụng', level=2)
doc.add_paragraph('Hệ thống được thiết kế phục vụ chính cho bộ phận Tuyển dụng và các nhóm liên quan trong quy trình đánh giá ứng viên.')

p = doc.add_paragraph(); bold_inline(p, 'HR / Chuyên viên Tuyển dụng (người dùng chính): ')
bullet(doc, '- Mục đích sử dụng: Nhập CV, JD, các link bài test và phỏng vấn, nhận báo cáo AI phân tích toàn diện.')
bullet(doc, '- Giá trị mang lại: Không cần đọc thủ công từng bài test. AI chấm điểm tự động và đưa ra phân tích có evidence cụ thể, HR chỉ cần review kết quả và ra quyết định tiếp theo.')

p = doc.add_paragraph(); bold_inline(p, 'Hiring Manager / HOD (người dùng phụ): ')
bullet(doc, '- Mục đích sử dụng: Xem báo cáo PDF tóm tắt năng lực ứng viên theo từng JD cụ thể trước khi phỏng vấn vòng cuối.')
bullet(doc, '- Giá trị mang lại: Được cung cấp data-driven insights về điểm mạnh, điểm yếu và mức độ phù hợp của ứng viên với vị trí và văn hóa AI-First.')

p = doc.add_paragraph(); bold_inline(p, 'Ban Giám đốc (người phê duyệt cuối): ')
bullet(doc, '- Mục đích sử dụng: Nhận báo cáo tóm tắt Decision (ĐẠT/KHÔNG ĐẠT) kèm AI Readiness Index và SWAT Elite Score trước khi phê duyệt offer.')
bullet(doc, '- Giá trị mang lại: Không cần đọc toàn bộ hồ sơ chi tiết. Chỉ cần xem báo cáo AI để ra quyết định nhanh theo tiêu chí NoAI-NoHire.')

doc.add_heading('1.3 Phạm vi chức năng', level=2)
doc.add_paragraph('Phạm vi hệ thống tập trung vào quy trình đánh giá ứng viên từ khi có hồ sơ đến khi ra quyết định:')
bullet(doc, '- Về loại hình dữ liệu: Hỗ trợ đọc và phân tích file CV (PDF/DOCX), nội dung JD text, link bài Test AI từ Odoo Survey, link bài Test 5G từ Odoo Survey, link báo cáo IQ/EQ từ jobtest.vn, link phỏng vấn Survey.')
bullet(doc, '- Về quy trình áp dụng: Ứng dụng sau vòng test năng lực và trước vòng phỏng vấn Hiring Manager, hoặc sau phỏng vấn trước khi trình Ban Giám đốc.')
bullet(doc, '- Về ranh giới hệ thống: AI chấm điểm và phân tích dựa trên dữ liệu có sẵn. Hệ thống KHÔNG tự ra quyết định offer – HR vẫn là người quyết định cuối cùng.')

p = doc.add_paragraph(); bold_inline(p, 'Các chức năng chính: Đánh Giá và Phân tích Ứng Viên ')
bullet(doc, '- Chấm điểm AI Test (10 tiêu chí): AI Awareness, AI Daily Use, AI Self-Assessment, AI Problem Solving, Prompt Engineering, AI×Teamwork, AI Productivity, AI Mindset, AI Limitation, AI Growth Plan – thang 100 điểm.')
bullet(doc, '- Chấm điểm 5G Test (5 nhóm G1–G5): Giao tiếp, Giao lưu thực tế, Giám sát, Giải quyết vấn đề, Giảng dạy – thang 100 điểm.')
bullet(doc, '- SWAT Elite Assessment (4 trụ cột): AI First Mindset (50%), 2AS Execution (20%), Practical Efficiency (20%), Risk Control & Language (10%) – thang 10 điểm, tự động ra Decision.')
bullet(doc, '- Phân tích định tính: Strengths / Gaps / Best At có evidence trích dẫn từng nguồn cụ thể.')

# Bảng chức năng
doc.add_paragraph()
tbl2 = doc.add_table(rows=5, cols=2)
tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
table_borders(tbl2, '4C94D8')
make_tbl_width(tbl2, 16)
func_data = [
    ('Chức năng', 'Mô tả'),
    ('Chấm điểm AI Test (10 tiêu chí)', 'Phân tích AI Readiness từ bài test Odoo Survey – thang 100 điểm, nhãn AI-Ready/Non-AI'),
    ('Chấm điểm 5G Test (G1–G5)', 'Đánh giá 5 nhóm tiêu chí tiềm năng quản lý từ Odoo Survey – thang 100 điểm'),
    ('SWAT Elite Assessment (4 trụ cột)', 'Chấm 4 trụ cột AI-First theo trọng số – thang 10 điểm, tự động Decision ĐẠT/KHÔNG ĐẠT'),
    ('Phân tích định tính & Xuất báo cáo', 'Strengths/Gaps/Best At có evidence + xuất báo cáo PDF A4 tự động trong 30–60 giây'),
]
for ri, (f, d) in enumerate(func_data):
    c0 = tbl2.rows[ri].cells[0]; c1 = tbl2.rows[ri].cells[1]
    c0.text = f; c1.text = d
    for c in [c0, c1]:
        runs = c.paragraphs[0].runs
        if runs: runs[0].font.size = Pt(11)
    if ri == 0:
        for c in [c0, c1]:
            if c.paragraphs[0].runs: c.paragraphs[0].runs[0].bold = True
            set_cell_shading(c, 'DCE6F1')

p = doc.add_paragraph()
r = p.add_run('Bảng 1: Chức năng chính của hệ thống DAIT AI Candidate Report')
r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x55,0x55,0x55)

doc.add_heading('1.4 Đường dẫn dự án', level=2)
p = doc.add_paragraph()
r = p.add_run('http://localhost:5200/assets/ai_ats/frontend/')
r.font.color.rgb = RGBColor(0x00, 0x56, 0xA0)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# CHƯƠNG 2: PHÂN QUYỀN HỆ THỐNG
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_heading('PHÂN QUYỀN HỆ THỐNG', level=1)

tbl3 = doc.add_table(rows=4, cols=2)
tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
table_borders(tbl3, '4C94D8')
make_tbl_width(tbl3, 16)
role_data = [
    ('Vai trò', 'Quyền hạn chính'),
    ('HR / Chuyên viên Tuyển dụng', 'Upload CV (PDF/DOCX), nhập JD text, nhập link các bài test và Survey, nhấn Generate, xem báo cáo AI, xuất PDF'),
    ('Hiring Manager / HOD', 'Xem báo cáo PDF tóm tắt, xem phân tích Strengths/Gaps/Best At theo JD cụ thể, đưa ra ý kiến phỏng vấn vòng cuối'),
    ('Ban Giám đốc', 'Nhận báo cáo PDF tóm tắt Decision (ĐẠT/KHÔNG ĐẠT) kèm AI Readiness Index + SWAT Elite Score, phê duyệt offer'),
]
for ri, (role, perm) in enumerate(role_data):
    c0 = tbl3.rows[ri].cells[0]; c1 = tbl3.rows[ri].cells[1]
    c0.text = role; c1.text = perm
    for c in [c0, c1]:
        if c.paragraphs[0].runs: c.paragraphs[0].runs[0].font.size = Pt(11)
    if ri == 0:
        for c in [c0, c1]:
            if c.paragraphs[0].runs: c.paragraphs[0].runs[0].bold = True
            set_cell_shading(c, 'DCE6F1')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════════
# CHƯƠNG 3: HƯỚNG DẪN SỬ DỤNG CHI TIẾT
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_heading('HƯỚNG DẪN SỬ DỤNG CHI TIẾT', level=1)

# 3.1
doc.add_heading('3.1 Truy cập và khởi tạo', level=2)

p = doc.add_paragraph(); bold_inline(p, 'Bước 1: ', 'Truy cập hệ thống qua đường dẫn nội bộ được cấp')
p = doc.add_paragraph(); bold_inline(p, 'Bước 2: ', 'Đăng nhập bằng tài khoản Frappe nội bộ.')
doc.add_paragraph('Hệ thống sẽ tự động:')
bullet(doc, '- Xác thực tài khoản và kiểm tra quyền truy cập app "ai_ats"')
bullet(doc, '- Hiển thị màn hình nhập liệu với ô upload CV, ô JD text, các ô nhập link test và nút Generate')
bullet(doc, '- Khởi tạo Session ID bảo mật cho phiên làm việc và ghi log hoạt động tự động')

doc.add_paragraph()
add_image_para(doc, SS/'fe_01_empty.png', 5.8, 'Ảnh 1: Màn hình giao diện chính – AI Candidate Evaluation Persona')

p = doc.add_paragraph(); bold_inline(p, 'Bước 3: ', 'Upload file CV (PDF hoặc DOCX) và điền nội dung JD vào ô tương ứng')
doc.add_paragraph()
add_image_para(doc, SS/'fe_02_filled.png', 5.8, 'Ảnh 2: Giao diện sau khi điền đầy đủ JD và Survey URL (bắt buộc)')

p = doc.add_paragraph(); bold_inline(p, 'Bước 4: ', 'Nhập các link bài test, sau đó nhấn nút "🚀 Generate Candidate Persona Report" và chờ kết quả (30–60 giây).')
bullet(doc, '- Odoo AI Test URL: Link bài Test AI từ Odoo Survey (hr-dev.ctgroupvietnam.com → In kết quả)')
bullet(doc, '- Odoo 5G Competency URL: Link bài Test 5G từ Odoo Survey')
bullet(doc, '- jobtest.vn EQ/IQ URL: Link báo cáo IQ/EQ từ jobtest.vn (nếu có)')
bullet(doc, '- Interview Survey URL (*bắt buộc*): Link phỏng vấn Survey của ứng viên')

doc.add_paragraph()
add_image_para(doc, SS/'fe_02b_button.png', 5.8, 'Ảnh 3: Giao diện nhập link bài test – nút Generate ở cuối trang')

# 3.2
doc.add_heading('3.2 Thực hiện phân tích ứng viên', level=2)

p = doc.add_paragraph(); bold_inline(p, 'Bước 1: ', 'Nhấn ô "Candidate Resume / CV" để upload file CV của ứng viên.')
bullet(doc, 'Tên file mẫu: BDLUV - IT TEST [Tên ứng viên] (full kn).pdf hoặc .docx')

p = doc.add_paragraph(); bold_inline(p, 'Bước 2: ', 'Paste nội dung JD vào ô "Job Description Content".')
bullet(doc, 'JD cần đủ: Tên vị trí, yêu cầu kỹ năng bắt buộc/phụ, mô tả công việc chính.')

p = doc.add_paragraph(); bold_inline(p, 'Bước 3: ', 'Điền các link bài test vào đúng ô tương ứng.')
bullet(doc, 'Link AI Test + 5G Test: Lấy từ Odoo Survey → mục Kết quả phỏng vấn ứng viên → nút In/Print.')
bullet(doc, 'Link IQ/EQ: Lấy từ jobtest.vn sau khi ứng viên hoàn thành bài test.')

p = doc.add_paragraph(); bold_inline(p, 'Bước 4: ', 'Nhấn nút "🚀 Generate Candidate Persona Report". Hệ thống tự động scrape, phân tích và trả về báo cáo.')

# 3.3
doc.add_heading('3.3 Đọc kết quả đánh giá', level=2)
doc.add_paragraph('Sau khi AI xử lý xong, hệ thống trả về đầy đủ các thông tin sau:')
bullet(doc, '- Decision tổng thể: ĐẠT (màu xanh) hoặc KHÔNG ĐẠT (màu đỏ), kèm lý do theo quy tắc NoAI-NoHire.')
bullet(doc, '- AI Readiness Index: Điểm AI Test /100 và nhãn phân loại [AI-Ready] hoặc [Non-AI].')
bullet(doc, '- SWAT Elite Score: Điểm SWAT /10 theo 4 trụ cột có trọng số – nhãn [Swat-Elite] hoặc [KHÔNG ĐẠT].')
bullet(doc, '- 5G Potential Score: Điểm 5G /100 theo 5 nhóm G1–G5 và nhãn phân loại (Giỏi/Trung bình/Không phù hợp).')
bullet(doc, '- Phân tích định tính: Strengths / Gaps / Best At với evidence trích dẫn cụ thể từng nguồn.')
doc.add_paragraph('Sau khi xem kết quả, nhấn nút "Xuất báo cáo PDF" để xuất báo cáo định dạng A4, sẵn sàng gửi Hiring Manager và Ban Giám đốc.')

doc.add_paragraph()
add_image_para(doc, SS/'ss_header.png', 5.8, 'Ảnh 4: Header báo cáo AI – Decision ĐẠT màu xanh và thông tin ứng viên')
doc.add_paragraph()
add_image_para(doc, SS/'ss_scores.png', 5.8, 'Ảnh 5: Bảng điểm tổng hợp – AI Readiness 84/100 [AI-Ready] · SWAT 8.0/10 [Swat-Elite] · 5G 86/100 [Giỏi]')
doc.add_paragraph()
add_image_para(doc, SS/'ss_section_1.png', 5.8, 'Ảnh 5b: Chi tiết bảng chấm điểm AI Test – 10 tiêu chí có progress bar và lý do')
doc.add_paragraph()
add_image_para(doc, SS/'ss_section_2.png', 5.8, 'Ảnh 5c: Bảng SWAT Elite Assessment – 4 trụ cột trọng số với cơ sở đánh giá')
doc.add_paragraph()
add_image_para(doc, SS/'ss_section_3.png', 5.8, 'Ảnh 5d: Bảng 5G Test – 5 nhóm tiêu chí G1–G5 tiềm năng quản lý')

# 3.4
doc.add_heading('3.4 Phân tích định tính & Tiện ích mở rộng', level=2)
doc.add_paragraph('Bên cạnh điểm số, hệ thống cung cấp phân tích định tính chi tiết giúp HR và Hiring Manager hiểu sâu về năng lực ứng viên:')

p = doc.add_paragraph(); bold_inline(p, 'Phần Strengths (Điểm mạnh):')
bullet(doc, '- Kỹ năng công nghệ: Liệt kê từng tool/tech với mức độ thành thạo và evidence từ ≥2 nguồn (CV + Test).')
bullet(doc, '- Năng lực vượt chuẩn JD: Những điểm ứng viên vượt trội so với yêu cầu JD, có thể tạo giá trị ngay.')
bullet(doc, '- Phù hợp JD: Từng yêu cầu cốt lõi được đánh giá ✓✓/✓/~ kèm evidence cụ thể.')

p = doc.add_paragraph(); bold_inline(p, 'Phần Gaps (Điểm hạn chế):')
bullet(doc, '- Kỹ năng/năng lực thiếu: Phân loại [THIẾU CỨNG], [CẦN CẢI THIỆN], [TƯ DUY AI-FIRST] kèm trích dẫn câu trả lời.')
bullet(doc, '- Rủi ro vận hành: Rủi ro bảo mật, hiệu suất, văn hóa AI và reliability dựa trên bài làm.')

p = doc.add_paragraph(); bold_inline(p, 'Phần Best At (Năng lực nổi bật):')
bullet(doc, '- Chuyên môn core + top 3 năng lực đóng góp ngay từ tuần đầu có evidence.')
bullet(doc, '- Năng lực vận hành 2AS tools: Đã dùng AI tool nào, mức hands-on, tiềm năng 30–60 ngày.')
bullet(doc, '- AI adoption level: [AI-Native/Willing/Hesitant/Resistant] kèm evidence từ test và CV.')

doc.add_paragraph()
add_image_para(doc, SS/'ss_section_4.png', 5.8, 'Ảnh 6: Phần Strengths – Điểm mạnh kỹ thuật và năng lực vượt chuẩn JD có evidence')
doc.add_paragraph()
add_image_para(doc, SS/'ss_section_5.png', 5.8, 'Ảnh 6b: Phần Gaps – Điểm hạn chế và rủi ro vận hành có evidence cụ thể')
doc.add_paragraph()
add_image_para(doc, SS/'ss_section_6.png', 5.8, 'Ảnh 6c: Phần Best At – Định vị năng lực nổi bật nhất và mức độ AI-readiness của ứng viên')

# 3.5
doc.add_heading('3.5 Kết quả phân tích mẫu thực tế', level=2)
doc.add_paragraph(
    'Dưới đây là báo cáo AI tự động xuất ra sau khi HR nhập đầy đủ: CV PDF + JD Full-Stack Developer (AI-First) + '
    'Link AI Test + Link 5G Test + Link IQ Test của ứng viên Trương Hoàng Vũ – '
    'dùng làm mẫu minh họa quy trình đánh giá thực tế.'
)

doc.add_paragraph()
add_image_para(doc, SS/'ss_full.png', 5.5, 'Ảnh 7: Báo cáo AI – Tổng quan (Header Decision ĐẠT + Thông tin ứng viên + 3 Score Cards)')
doc.add_paragraph()
add_image_para(doc, SS/'ss_analysis.png', 5.5, 'Ảnh 8: Báo cáo AI – Chi tiết phân tích (Bảng AI Test, SWAT, 5G và nhận xét định tính)')

# ── Lưu file ──────────────────────────────────────────────────────────────────
doc.save(str(OUTPUT))
print(f"\n✅ Tài liệu hoàn chỉnh đã tạo!")
print(f"   → {OUTPUT}")
print(f"   Size: {OUTPUT.stat().st_size / 1024:.0f} KB")
